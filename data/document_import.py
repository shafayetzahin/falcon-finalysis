"""Review-first extraction of financial statement rows from PDF and XLSX files."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
import zipfile

import pandas as pd
from openpyxl import load_workbook

from core.config import ALIASES, FIELDS
from data.parsers import normalize


YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")


@dataclass
class Extraction:
    frame: pd.DataFrame
    evidence: pd.DataFrame
    unit_hint: str
    notes: list[str]


def _field_lookup() -> dict[str, str]:
    lookup = {normalize(field): field for field in FIELDS}
    for target, aliases in ALIASES.items():
        if target != "Year":
            lookup.update({normalize(alias): target for alias in aliases})
    return lookup


LOOKUP = _field_lookup()


def _field(label: object) -> str | None:
    clean = normalize(str(label))
    if clean in LOOKUP:
        return LOOKUP[clean]
    # Annual reports often add numbering, notes and qualifiers to the row name.
    matches = [(len(alias), target) for alias, target in LOOKUP.items()
               if len(alias) >= 4 and (alias in clean or clean in alias)]
    return max(matches, default=(0, None))[1]


def _amount(value: object) -> float | None:
    if value is None or pd.isna(value):
        return None
    text = str(value).strip().replace(",", "").replace("৳", "").replace("$", "")
    if text in {"", "-", "—", "–", "n/a", "N/A"}:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return float(text) if text not in {"", "-", "."} else None
    except ValueError:
        return None


def _extract_rows(rows: list[list[object]], source: str) -> tuple[list[dict], list[dict]]:
    records: dict[int, dict] = {}
    evidence: list[dict] = []
    for header_i, row in enumerate(rows):
        positions: list[tuple[int, int]] = []
        for col_i, value in enumerate(row):
            match = YEAR_RE.search(str(value or ""))
            if match:
                positions.append((col_i, int(match.group())))
        if not positions:
            continue
        for data_row in rows[header_i + 1:header_i + 55]:
            label = next((x for x in data_row[:max(p[0] for p in positions)] if str(x or "").strip()), None)
            target = _field(label) if label is not None else None
            if not target:
                continue
            for col_i, year in positions:
                if col_i >= len(data_row):
                    continue
                value = _amount(data_row[col_i])
                if value is None:
                    continue
                records.setdefault(year, {"Year": year})[target] = value
                evidence.append({"Year": year, "Falcon Finalysis field": target,
                                 "Source label": str(label).strip(), "Value": value, "Source": source})
        if records:
            break
    return list(records.values()), evidence


def _extract_text_rows(text: str, source: str) -> tuple[list[dict], list[dict]]:
    """Fallback for visually aligned PDF tables without drawn cell borders."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]
    for header_i, line in enumerate(lines):
        years = [int(x) for x in YEAR_RE.findall(line)]
        if not years:
            continue
        records = {year: {"Year": year} for year in years}
        evidence = []
        for item in lines[header_i + 1:header_i + 55]:
            tokens = item.split()
            if len(tokens) <= len(years):
                continue
            values = [_amount(x) for x in tokens[-len(years):]]
            if not all(value is not None for value in values):
                continue
            label = " ".join(tokens[:-len(years)])
            target = _field(label)
            if not target:
                continue
            for year, value in zip(years, values):
                records[year][target] = value
                evidence.append({"Year": year, "Falcon Finalysis field": target, "Source label": label,
                                 "Value": value, "Source": source})
        if evidence:
            return list(records.values()), evidence
    return [], []


def _combine(parts: list[Extraction]) -> Extraction:
    if not parts:
        raise ValueError("No recognizable financial statement rows were found.")
    all_rows = pd.concat([p.frame for p in parts], ignore_index=True)
    rows = []
    conflicts = 0
    for year, group in all_rows.groupby("Year"):
        row = {"Year": int(year)}
        for field in FIELDS:
            values = group[field].dropna().unique() if field in group else []
            if len(values):
                row[field] = values[0]
                conflicts += max(0, len(values) - 1)
        rows.append(row)
    frame = pd.DataFrame(rows).reindex(columns=["Year"] + FIELDS).sort_values("Year").reset_index(drop=True)
    evidence = pd.concat([p.evidence for p in parts], ignore_index=True)
    notes = [n for p in parts for n in p.notes]
    if conflicts:
        notes.append(f"{conflicts} duplicate year/field value(s) differed; the first extracted value is shown. Review before import.")
    hints = [p.unit_hint for p in parts if p.unit_hint != "Unknown"]
    return Extraction(frame, evidence, hints[0] if len(set(hints)) == 1 else "Unknown", notes)


def extract_pdfs(files: list[tuple[bytes, str]]) -> Extraction:
    """Extract candidates from text-based PDF tables; values always require review."""
    import pdfplumber

    parts: list[Extraction] = []
    for payload, filename in files:
        if not payload or len(payload) > 40 * 1024 * 1024:
            raise ValueError(f"{filename}: use a non-empty PDF smaller than 40 MB.")
        try:
            with pdfplumber.open(BytesIO(payload)) as pdf:
                if len(pdf.pages) > 300:
                    raise ValueError(f"{filename}: PDFs are limited to 300 pages.")
                records, evidence, hints = [], [], []
                for page_no, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    lower = text.lower()
                    if "in million" in lower or "million taka" in lower or "bdt million" in lower:
                        hints.append("Millions")
                    elif "in thousand" in lower or "thousand taka" in lower or "bdt thousand" in lower:
                        hints.append("Thousands")
                    if not any(term in lower for term in ["revenue", "profit", "assets", "liabilities", "cash flow", "financial position", "income statement"]):
                        continue
                    for table_no, table in enumerate(page.extract_tables() or [], 1):
                        found, proof = _extract_rows(table, f"{filename}, page {page_no}, table {table_no}")
                        records.extend(found)
                        evidence.extend(proof)
                    if not any(str(item.get("Source", "")).startswith(f"{filename}, page {page_no},") for item in evidence):
                        found, proof = _extract_text_rows(text, f"{filename}, page {page_no}, aligned text")
                        records.extend(found)
                        evidence.extend(proof)
                if records:
                    frame = pd.DataFrame(records).reindex(columns=["Year"] + FIELDS)
                    parts.append(Extraction(frame, pd.DataFrame(evidence),
                                            hints[0] if len(set(hints)) == 1 else "Unknown",
                                            [f"{filename}: extracted from text tables; scanned-image pages need OCR before upload."]))
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"{filename}: unable to read this PDF. Try a text-based annual report or export its statements to Excel.") from exc
    if not parts:
        raise ValueError("No recognizable statement tables were found. Scanned PDFs need OCR; otherwise use Excel or the manual editor.")
    return _combine(parts)


def extract_workbook(payload: bytes, filename: str) -> Extraction:
    """Find metric-by-year financial statement layouts across values-only sheets."""
    if not payload or len(payload) > 15 * 1024 * 1024 or Path(filename).suffix.lower() != ".xlsx":
        raise ValueError("Use a values-only .xlsx file smaller than 15 MB.")
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 75 * 1024 * 1024:
                raise ValueError("The expanded workbook is too large (75 MB limit).")
        book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        parts = []
        try:
            for ws in book.worksheets:
                if ws.max_row > 500 or ws.max_column > 100:
                    continue
                rows = []
                for row in ws.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("Formula cells are not accepted. Paste values into a copy before upload.")
                    rows.append([cell.value for cell in row])
                records, evidence = _extract_rows(rows, f"{filename}, sheet {ws.title}")
                if records:
                    parts.append(Extraction(pd.DataFrame(records).reindex(columns=["Year"] + FIELDS),
                                            pd.DataFrame(evidence), "Unknown",
                                            [f"Read statement-style rows from worksheet '{ws.title}'."]))
        finally:
            book.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read this workbook. Check that it is a valid values-only .xlsx file.") from exc
    if not parts:
        raise ValueError("No metric-by-year statement layout was found. Use the standard column mapping importer instead.")
    return _combine(parts)


def apply_scale(frame: pd.DataFrame, multiplier: float) -> pd.DataFrame:
    out = frame.copy()
    out[FIELDS] = out[FIELDS].apply(pd.to_numeric, errors="coerce") * float(multiplier)
    return out
