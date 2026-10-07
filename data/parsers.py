"""Bounded CSV/XLSX input and configurable financial-label mapping."""
from io import BytesIO
import csv
from numbers import Number
from pathlib import Path
import re
import zipfile
import pandas as pd
from openpyxl import load_workbook
from core.config import ALIASES, FIELDS
from core.validation import validate


def normalize(label: str) -> str:
    return re.sub(r'[^a-z0-9]', '', str(label).lower())


def parse_date(value: object) -> pd.Timestamp:
    """Daily dates use day-first order, except explicit year-first ISO dates."""
    if isinstance(value, Number):
        return pd.NaT
    text = str(value).strip()
    if text.isdigit():
        return pd.NaT
    iso = bool(re.match(r'^\d{4}[-/]\d{2}[-/]\d{2}(?:$|[ T])', text))
    try:
        parsed = pd.to_datetime(value, errors='coerce', dayfirst=not iso)
        if pd.isna(parsed):
            return pd.NaT
        return pd.Timestamp(parsed).tz_localize(None).normalize()
    except (TypeError, ValueError, OverflowError):
        return pd.NaT


def numeric_values(series: pd.Series) -> pd.Series:
    """Parse whole numeric cells, including scientific notation and accounting negatives."""
    text = series.astype('string').str.strip()
    text = text.str.replace(r'^\((.*)\)$', r'-\1', regex=True)
    # Accept western or South Asian grouping, never turn a malformed 1,2 into 12.
    grouped = (r'[-+]?(?:\d{1,3}(?:,\d{3})+|\d{1,2}(?:,\d{2})*,\d{3})'
               r'(?:\.\d+)?(?:[eE][-+]?\d+)?')
    malformed = text.str.contains(',', regex=False, na=False) & ~text.str.fullmatch(grouped, na=False)
    text = text.mask(malformed).str.replace(',', '', regex=False)
    return pd.to_numeric(text, errors='coerce')


def suggest_mapping(columns: list, extra: dict | None = None) -> dict[str, str]:
    lookup = {normalize(c): c for c in ['Year'] + FIELDS}
    for target, aliases in ALIASES.items():
        lookup.update({normalize(a): target for a in aliases})
    lookup.update({normalize(k): v for k, v in (extra or {}).items()})
    return {str(c): lookup.get(normalize(c), 'Ignore') for c in columns}


def read_tabular(payload: bytes, filename: str, sheet: str | None = None, *,
                 max_rows: int = 100, max_columns: int = 100,
                 max_bytes: int = 10 * 1024 * 1024) -> pd.DataFrame:
    """Bound tabular imports and reject duplicate headers and cached Excel formulas."""
    if not payload:
        raise ValueError('The uploaded file is empty.')
    if len(payload) > max_bytes:
        raise ValueError(f'Upload a file smaller than {max_bytes // (1024 * 1024)} MB.')
    def check_headers(headers):
        labels = [str(value).strip() if value is not None else '' for value in headers]
        if not labels or any(not value for value in labels) or len(labels) != len(set(labels)):
            raise ValueError('Use unique, non-empty column headers.')
        if len(labels) > max_columns:
            raise ValueError(f'Use at most {max_columns:,} columns.')
    size_error = f'Use at most {max_rows:,} rows and {max_columns:,} columns.'
    ext = Path(filename).suffix.lower()
    try:
        if ext == '.csv':
            header = next(csv.reader(payload.decode('utf-8-sig').splitlines()), [])
            check_headers(header)
            df = pd.read_csv(BytesIO(payload), nrows=max_rows + 1, encoding='utf-8-sig')
        elif ext == '.xlsx':
            with zipfile.ZipFile(BytesIO(payload)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('The expanded workbook is too large (50 MB limit).')
            book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
            try:
                ws = book[sheet] if sheet else book.active
                if (ws.max_row and ws.max_row > max_rows + 1
                        or ws.max_column and ws.max_column > max_columns):
                    raise ValueError(size_error)
                # Dimensions are supplied by the file and can understate its real size.
                ws.reset_dimensions()
                rows = []
                for row in ws.iter_rows():
                    if len(rows) > max_rows or len(row) > max_columns:
                        raise ValueError(size_error)
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Formula cells are not accepted. Paste values into a copy before upload.')
                    rows.append([cell.value for cell in row])
                if not rows:
                    raise ValueError('This worksheet is empty.')
                width = max(len(row) for row in rows)
                headers = rows[0] + [None] * (width - len(rows[0]))
                check_headers(headers)
                df = pd.DataFrame(rows[1:], columns=headers).dropna(how='all')
            finally:
                book.close()
        else:
            raise ValueError('Use .xlsx or UTF-8 .csv. Convert legacy .xls files to .xlsx first.')
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Unable to read this file. Check its format and use the downloadable template.') from exc
    if df.empty:
        raise ValueError('The file contains no data rows.')
    if len(df) > max_rows or len(df.columns) > max_columns:
        raise ValueError(size_error)
    return df


def read_file(payload: bytes, filename: str, sheet: str | None = None) -> pd.DataFrame:
    """Read a compact financial input table."""
    return read_tabular(payload, filename, sheet)


def prepare(raw: pd.DataFrame, mapping: dict | None = None) -> pd.DataFrame:
    """Apply user-reviewed mappings; do not infer subtotals or replace missing amounts."""
    mapping = mapping or suggest_mapping(list(raw.columns))
    selected = {k: v for k, v in mapping.items() if v != 'Ignore'}
    if len(selected.values()) != len(set(selected.values())):
        raise ValueError('Two columns map to the same field. Assign one or ignore the duplicate.')
    if any(v not in ['Year'] + FIELDS for v in selected.values()):
        raise ValueError('Mapping contains an unknown financial field.')
    if any(column not in raw.columns for column in selected):
        raise ValueError('The mapping refers to a column that is not present in this file.')
    df = raw[list(selected)].rename(columns=selected).copy()
    for col in df:
        # Commas and accounting parentheses are common values-only export conventions.
        text = df[col].astype('string').str.strip()
        numeric = numeric_values(df[col])
        if (text.notna() & text.ne('') & numeric.isna()).any():
            raise ValueError(f'{col} contains text or formulas. Use numbers, or blank for unavailable data.')
        df[col] = numeric.astype(float)
    errors = [x.message for x in validate(df) if x.severity == 'ERROR']
    if errors:
        raise ValueError(' '.join(errors))
    df['Year'] = df['Year'].astype(int)
    return df.reindex(columns=['Year'] + FIELDS).sort_values('Year').reset_index(drop=True)
