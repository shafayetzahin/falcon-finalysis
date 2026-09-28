"""Styled, values-only Excel snapshots; no formulas or active content are exported."""
from io import BytesIO
import math
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from core.config import FIELDS, INCOME, BALANCE, CASH_FLOW, OPTIONAL, DISCLAIMER, AVERAGE_POLICY
from core.ratio_engine import BY_NAME, METRICS
from core.trend_engine import horizontal
from core.scenario_engine import Scenario, baseline, model, ASSUMPTIONS

NAVY = '173859'
TEAL = '008C95'
ACTIVE_LOGO = Path(__file__).resolve().parents[1] / 'assets' / 'logo-active.png'


def market_data_workbook(history: pd.DataFrame, exchange: str, ticker: str, source_url: str,
                         fetched_at: str) -> bytes:
    """Create a clean, typed workbook for an official exchange price extract."""
    if history.empty:
        raise ValueError('Price history is empty.')
    data = history.copy().sort_values('Date').reset_index(drop=True)
    data['Date'] = pd.to_datetime(data['Date'])
    close = pd.to_numeric(data['Close'], errors='coerce')
    data['Daily Return'] = close.pct_change(fill_method=None)
    monthly = (data.set_index('Date').resample('ME')
               .agg(Close=('Close', 'last'), Average_Close=('Close', 'mean'),
                    Trading_Days=('Close', 'count'), Volume=('Volume', 'sum'))
               .dropna(subset=['Close']).reset_index())

    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.sheet_view.showGridLines = False
    summary.append([f'{exchange} market history - {ticker}'])
    summary.append(['Period', f'{data.Date.min():%d %b %Y} to {data.Date.max():%d %b %Y}'])
    summary.append(['Trading records', len(data)])
    summary.append(['First close', float(close.iloc[0])])
    summary.append(['Last close', float(close.iloc[-1])])
    summary.append(['Period return', float(close.iloc[-1] / close.iloc[0] - 1) if close.iloc[0] else None])
    summary.append(['Total volume', float(pd.to_numeric(data.get('Volume'), errors='coerce').sum())])
    summary.append([])
    summary.append(['Use', 'Official exchange values arranged for filtering, charting and further analysis.'])
    summary.append(['Source', source_url])
    summary.append(['Fetched at (UTC)', fetched_at])
    summary.append(['Caution', 'Exchange website data may be delayed, corrected, unavailable or reformatted. Verify material decisions at the source.'])
    summary['A1'].font = Font(size=16, bold=True, color=NAVY)
    for cell in summary['A']:
        if cell.row > 1:
            cell.font = Font(bold=True, color=NAVY)
    summary['B6'].number_format = '0.00%'
    summary.column_dimensions['A'].width = 23
    summary.column_dimensions['B'].width = 90
    summary.freeze_panes = 'A2'

    prices = book.create_sheet('Price History')
    prices.sheet_view.showGridLines = False
    headers = list(data.columns)
    prices.append(headers)
    for row in data.itertuples(index=False, name=None):
        prices.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else safe_cell(v) for v in row])
    prices.freeze_panes = 'A2'
    prices.auto_filter.ref = prices.dimensions
    for cell in prices[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center')
    for cell in prices['A'][1:]:
        cell.number_format = 'dd-mmm-yyyy'
    if 'Daily Return' in headers:
        ret_col = headers.index('Daily Return') + 1
        for cell in list(prices.columns)[ret_col - 1][1:]:
            cell.number_format = '0.00%'
    for idx, name in enumerate(headers, 1):
        prices.column_dimensions[get_column_letter(idx)].width = max(12, min(22, len(name) + 3))
    chart = LineChart()
    chart.title = f'{ticker} closing price'
    chart.y_axis.title = 'Price (BDT)'
    chart.x_axis.title = 'Date'
    close_col = headers.index('Close') + 1
    chart.add_data(Reference(prices, min_col=close_col, min_row=1, max_row=prices.max_row), titles_from_data=True)
    chart.set_categories(Reference(prices, min_col=1, min_row=2, max_row=prices.max_row))
    chart.width, chart.height = 22, 9
    summary.add_chart(chart, 'D2')

    month = book.create_sheet('Monthly Summary')
    month.sheet_view.showGridLines = False
    month_headers = ['Month', 'Month-end Close', 'Average Close', 'Trading Days', 'Volume']
    month.append(month_headers)
    for row in monthly.itertuples(index=False, name=None):
        month.append(list(row))
    month.freeze_panes = 'A2'
    for cell in month[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
    for cell in month['A'][1:]:
        cell.number_format = 'mmm-yyyy'
    for col in range(1, 6):
        month.column_dimensions[get_column_letter(col)].width = 19

    sources = book.create_sheet('Sources')
    sources.sheet_view.showGridLines = False
    sources.append(['Item', 'Details'])
    sources.append(['Exchange', exchange])
    sources.append(['Ticker', ticker])
    sources.append(['Official source', source_url])
    sources.append(['Fetched at (UTC)', fetched_at])
    sources.append(['Method', 'Public exchange web page parsed by Falcon Finalysis; no price values are estimated.'])
    sources.append(['Refresh', 'Return to Listed Company Data and fetch the required date range again.'])
    for cell in sources[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
    sources.column_dimensions['A'].width = 22
    sources.column_dimensions['B'].width = 110
    sources.freeze_panes = 'A2'

    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def safe_cell(value):
    """Neutralize Excel formula injection in user-provided metadata."""
    if isinstance(value, str) and value[:1] in ['=', '+', '-', '@']:
        return "'" + value
    if value is None or not isinstance(value, (str, int, float, bool)) and pd.isna(value):
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    return value


def add_sheet(book: Workbook, title: str, df: pd.DataFrame, note: str = '', index: bool = True,
              percent: bool = False) -> None:
    ws = book.create_sheet(title)
    ws.append([title])
    ws.append([safe_cell(note)])
    headers = ([df.index.name or 'Metric'] if index else []) + [str(c) for c in df.columns]
    ws.append(headers)
    for idx, row in df.iterrows():
        values = ([idx] if index else []) + list(row)
        ws.append([safe_cell(v) for v in values])
    last = max(2, len(headers))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    ws['A1'].font = Font(name='Calibri', size=20, color='FFFFFF', bold=True)
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 65 if len(note) > 160 else 36
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='center')
    ws['A2'].font = Font(name='Calibri', size=10, color='587089')
    for cell in ws[3]:
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.font = Font(name='Calibri', color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[3].height = 32
    for row in ws.iter_rows(min_row=4):
        metric = str(row[0].value)
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color=NAVY)
            cell.fill = PatternFill('solid', fgColor='F1F5F9' if cell.row % 2 == 0 else 'FFFFFF')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(cell.value, (int, float)):
                is_pct = percent or title == 'Ratios' and metric in BY_NAME and BY_NAME[metric].unit == '%'
                cell.number_format = '0.0%;[Red](0.0%);0.0%' if is_pct else '#,##0.00;[Red](#,##0.00);0.00'
        longest = max((len(str(c.value or '')) for c in row if isinstance(c.value, str)), default=0)
        ws.row_dimensions[row[0].row].height = min(180, max(25, (longest//65+1)*15))
    for i in range(1, last+1):
        ws.column_dimensions[get_column_letter(i)].width = 36 if i == 1 else 22
    ws.freeze_panes = 'B4' if index else 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(last)}{ws.max_row}'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:3'


def input_template(frame: pd.DataFrame | None = None) -> bytes:
    """Values-only importable first sheet, with schema guidance on a separate sheet."""
    book = Workbook()
    ws = book.active
    ws.title = 'Financial Data'
    df = frame if frame is not None else pd.DataFrame({'Year': [2023, 2024, 2025]}).reindex(columns=['Year']+FIELDS)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([safe_cell(v) for v in row])
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for i in range(1, len(df.columns)+1):
        ws.column_dimensions[get_column_letter(i)].width = 23
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = '#,##0.00'
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = 'B2'
    guide = book.create_sheet('Field Guide')
    guide.append(['Field', 'Statement area', 'Priority', 'What to enter'])
    sections = {**{x: 'Income statement' for x in INCOME}, **{x: 'Balance sheet' for x in BALANCE},
                **{x: 'Cash flow' for x in CASH_FLOW}, **{x: 'Optional / market' for x in OPTIONAL}}
    for field in FIELDS:
        priority = 'Core analysis' if field not in OPTIONAL else 'Optional / advanced metrics'
        convention = ('Positive magnitude' if field in ['COGS', 'Operating Expenses', 'Depreciation',
                      'Interest Expense', 'Tax Expense', 'Capital Expenditure', 'Cash Dividends']
                      else 'Signed cash flow' if field in ['Operating Cash Flow', 'Investing Cash Flow',
                                                           'Financing Cash Flow', 'FX Effect on Cash']
                      else 'Reported value; blank if unavailable')
        guide.append([field, sections[field], priority, convention])
    for cell in guide[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center')
    guide.freeze_panes = 'A2'
    guide.auto_filter.ref = guide.dimensions
    guide.sheet_view.showGridLines = False
    for col, width in zip('ABCD', [35, 22, 28, 44]):
        guide.column_dimensions[col].width = width
    for row in guide.iter_rows(min_row=2):
        if row[2].value == 'Optional / advanced metrics':
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='F3F6F9')
    instructions = book.create_sheet('Instructions')
    for line in ['Falcon Finalysis input template', 'Start with the Core analysis fields in the Field Guide. Optional fields unlock per-share and market ratios.',
                 'One row per full fiscal year, 3–10 consecutive years. All amounts in full currency units.',
                 'Expenses, depreciation, interest, capex and dividends: positive magnitudes. Cash flows: signed.',
                 'Operating Expenses excludes depreciation. Supply reported subtotals; Falcon Finalysis never silently derives them.',
                 'Blank means unavailable. Enter explicit zero where applicable, including Preferred Dividends.',
                 'EPS uses Weighted Average Shares Outstanding. Ending Shares Outstanding is used for book value and dividend yield.',
                 'Paste values only. Formulas and macros are not accepted. The first sheet is importable.',
                 'Demo workbook only: fictional Apex Consumer Industries, generated by Falcon Finalysis. Not actual company data.',
                 DISCLAIMER]:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 115
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True)
        instructions.row_dimensions[row[0].row].height = 34
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def exchange_input_template(frame: pd.DataFrame, details: pd.DataFrame, exchange: str,
                            ticker: str, source_url: str, fetched_at: str) -> bytes:
    """Add exchange-reported annual metrics to the canonical financial template."""
    book = load_workbook(BytesIO(input_template(frame)))
    ws = book.create_sheet('Exchange Data', 1)
    ws.append([f'{exchange} annual values - {ticker}'])
    ws.append(['Only directly reported values are copied to Financial Data. Complete remaining fields from annual reports.'])
    ws.append(['Year', 'Exchange metric', 'Value', 'Unit', 'Source'])
    for row in details.itertuples(index=False, name=None):
        ws.append([safe_cell(value) for value in row])
    ws.merge_cells('A1:E1')
    ws.merge_cells('A2:E2')
    ws['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
    ws['A2'].alignment = Alignment(wrap_text=True)
    for cell in ws[3]:
        cell.fill = PatternFill('solid', fgColor=TEAL)
        cell.font = Font(color='FFFFFF', bold=True)
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:E{ws.max_row}'
    ws.sheet_view.showGridLines = False
    for column, width in zip('ABCDE', [12, 26, 18, 16, 72]):
        ws.column_dimensions[column].width = width
    sources = book.create_sheet('Exchange Source', 2)
    for row in [('Item', 'Details'), ('Exchange', exchange), ('Ticker', ticker),
                ('Official source', source_url), ('Fetched at (UTC)', fetched_at),
                ('Copied to Financial Data', 'Net Income when the exchange reports annual net profit in BDT million.'),
                ('Still needed', 'Use annual reports to complete revenue, balance sheet, cash flow and other missing fields.')]:
        sources.append(list(row))
    for cell in sources[1]:
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color='FFFFFF', bold=True)
    sources.column_dimensions['A'].width = 26
    sources.column_dimensions['B'].width = 100
    sources.sheet_view.showGridLines = False
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def excel_report(a, meta: dict, scenario: dict | None = None,
                 provenance: pd.DataFrame | None = None) -> bytes:
    """Export all analyses as labelled snapshots with numeric types and appropriate formats."""
    book = Workbook()
    book.remove(book.active)
    note = f'{meta["company_name"]} · {meta["currency"]} · FY{a.statements.index.min()}–FY{a.statements.index.max()} · Values-only snapshot; rerun Falcon Finalysis after input changes.'
    add_sheet(book, 'Executive Summary', pd.DataFrame({'Interpretation': a.summary}), note)
    book['Executive Summary'].column_dimensions['B'].width = 110
    logo = XLImage(str(ACTIVE_LOGO))
    logo.width, logo.height = 235, 134
    book['Executive Summary'].add_image(logo, 'C1')
    add_sheet(book, 'Raw Data', a.statements.reset_index(), note, index=False)
    for name, cols in [('Income Statement', INCOME), ('Balance Sheet', BALANCE), ('Cash Flow', CASH_FLOW)]:
        add_sheet(book, name, a.statements[cols].T, note)
    add_sheet(book, 'Ratios', a.ratios.T, AVERAGE_POLICY)
    add_sheet(book, 'Growth', a.growth.T, 'YoY requires a positive prior base. Rates are unavailable for missing or nonpositive bases.', percent=True)
    add_sheet(book, 'CAGR', a.cagr.to_frame('CAGR'), 'Positive endpoints; at least three observations; elapsed fiscal years.', percent=True)
    add_sheet(book, 'Horizontal Analysis', horizontal(a.statements, int(a.statements.index[0])), 'Base year = first available year.', index=False)
    for row in book['Horizontal Analysis'].iter_rows(min_row=4):
        row[-1].number_format = '0.0%'
    add_sheet(book, 'Vertical Analysis', a.vertical.T, 'Income items / revenue; balance items / total assets.', percent=True)
    add_sheet(book, 'DuPont', a.dupont.T, 'Net margin × asset turnover × equity multiplier. ROE and net margin are decimal ratios.')
    for row in book['DuPont'].iter_rows(min_row=4):
        if row[0].value in ['Net Profit Margin', 'DuPont ROE']:
            for cell in row[1:]:
                cell.number_format = '0.0%'
    add_sheet(book, 'Working Capital', a.ratios[['Net Working Capital', 'Current Ratio', 'Inventory Days', 'Receivable Days', 'Payable Days', 'Cash Conversion Cycle']].T, AVERAGE_POLICY)
    add_sheet(book, 'Risk Flags', pd.DataFrame([f.to_dict() for f in a.flags]), 'Generic rules are investigation prompts, not advice.', index=False)
    for col in ['F', 'G']:
        book['Risk Flags'].column_dimensions[col].width = 65
    add_sheet(book, 'Health Score', a.health[-1].categories, f'Falcon Finalysis Financial Health Score: {a.health[-1].score if a.health[-1].score is not None else "N/A"}; {a.health[-1].label}. Not an industry-standard credit score.', index=False)
    try:
        base = baseline(a.statements.iloc[-1])
        case = Scenario(**scenario) if scenario else base
        add_sheet(book, 'Scenario Analysis', pd.DataFrame({'Base case': model(a.statements.iloc[-1], base), 'Scenario': model(a.statements.iloc[-1], case)}), 'One-year operating model; ROA/ROE use frozen denominators; cash flows are proxies.')
        add_sheet(book, 'Scenario Assumptions', pd.DataFrame({'Base case': base.to_dict(), 'Scenario': case.to_dict()}), ASSUMPTIONS)
    except ValueError as exc:
        add_sheet(book, 'Scenario Analysis', pd.DataFrame({'Status': [str(exc)]}), 'Scenario unavailable: complete the required inputs.', index=False)
    add_sheet(book, 'Data Quality', pd.DataFrame([vars(i) for i in a.issues], columns=['severity', 'year', 'message']), 'Warnings do not modify reported figures.', index=False)
    if isinstance(provenance, pd.DataFrame) and not provenance.empty:
        add_sheet(book, 'Data Sources', provenance, 'Cell-level source records captured during import or manual editing.', index=False)
    add_sheet(book, 'Methodology', pd.DataFrame({'Formula': {m.name: m.formula for m in METRICS}}), DISCLAIMER)
    book['Methodology'].column_dimensions['B'].width = 95
    chart = LineChart()
    chart.title = 'Revenue and net income'
    chart.y_axis.title = meta['currency'] + ' (full units)'
    ws = book['Income Statement']
    for r in [4, 14]:
        chart.add_data(Reference(ws, min_row=r, max_row=r, min_col=1, max_col=len(a.statements)+1), from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ws, min_row=3, max_row=3, min_col=2, max_col=len(a.statements)+1))
    chart.width, chart.height = 23, 10
    ws.add_chart(chart, 'A18')
    # Fit wrapped text after all specialized column widths are known.
    for sheet in book:
        widths = {i: sheet.column_dimensions[get_column_letter(i)].width for i in range(1, sheet.max_column+1)}
        note_text = str(sheet['A2'].value or '')
        sheet.row_dimensions[2].height = min(400, max(36, 14*math.ceil(len(note_text)/max(sum(widths.values())*.9, 1))+10))
        for row in sheet.iter_rows(min_row=4):
            lines = max((math.ceil(len(str(cell.value or ''))/max(widths[cell.column]*.85, 1))
                         for cell in row if isinstance(cell.value, str)), default=1)
            sheet.row_dimensions[row[0].row].height = min(400, max(25, lines*15+5))
        if sheet.title in ['Raw Data', 'Horizontal Analysis']:
            for row in sheet.iter_rows(min_row=4):
                row[0].number_format = '0'
        if sheet.title in ['Income Statement', 'Balance Sheet', 'Cash Flow', 'Raw Data']:
            code = str(meta['currency']).replace('"', '')
            for row in sheet.iter_rows(min_row=4, min_col=2):
                for cell in row:
                    if isinstance(cell.value, (int, float)):
                        if sheet.title == 'Raw Data' and sheet.cell(3, cell.column).value in ['Shares Outstanding', 'Weighted Average Shares Outstanding']:
                            cell.number_format = '#,##0'
                            continue
                        cell.number_format = f'"{code} "#,##0.00;[Red](#,##0.00);"{code} "0.00'
        if sheet.title == 'Health Score':
            for row in sheet.iter_rows(min_row=4):
                row[-1].number_format = '0.0%'
        if sheet.title == 'Scenario Analysis':
            for row in sheet.iter_rows(min_row=4):
                if row[0].value in ['Net Margin', 'ROA', 'ROE']:
                    for cell in row[1:]:
                        cell.number_format = '0.0%'
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
