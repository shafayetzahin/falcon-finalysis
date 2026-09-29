"""Export structure, round trips and Streamlit page/workflow smoke checks."""
from datetime import date
from io import BytesIO
from pathlib import Path
import ssl
import pytest
import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from core.financial_engine import analyze
from data.demo import demo_company, DEMO_META
from data.parsers import read_file, prepare
from reports.excel_report import input_template, excel_report, exchange_input_template, market_data_workbook
from reports.pdf_report import pdf_report
from data.document_import import extract_workbook, extract_pdfs, apply_scale
from data.market_data import (_annual_metrics, _standardize_dse, read_price_file,
                              ticker_list, bundled_ticker_list, bundled_ticker_catalog,
                              price_history, _verified_exchange_session, _modern_dse_history)


def test_ticker_suggestions_support_cse_suffixes_and_bundled_fallback(monkeypatch):
    class Response:
        text = "<option value=''>Choose</option><option value='AAMRANET|Main'>AAMRANET</option><option value='SQURPHARMA|Main'>SQURPHARMA</option>"

    monkeypatch.setattr('data.market_data._request', lambda *args, **kwargs: Response())
    assert ticker_list('CSE') == ['AAMRANET', 'SQURPHARMA']
    assert 'SQURPHARMA' in bundled_ticker_list('DSE')
    assert 'SQURPHARMA' in bundled_ticker_list('CSE')
    assert 'PHARMACEUTICAL' in bundled_ticker_catalog('CSE')['SQURPHARMA']

ROOT = Path(__file__).resolve().parents[1]


def test_xlsx_roundtrip_and_export():
    demo = demo_company()
    restored = prepare(read_file(input_template(demo), 'demo.xlsx'))
    assert restored['Revenue'].tolist() == demo['Revenue'].tolist()
    a = analyze(restored)
    data = excel_report(a, {**DEMO_META, 'company_name': '=HYPERLINK("bad")'})
    book = load_workbook(BytesIO(data), data_only=False)
    required = ['Executive Summary', 'Raw Data', 'Income Statement', 'Balance Sheet', 'Cash Flow',
                'Ratios', 'Growth', 'Horizontal Analysis', 'Vertical Analysis', 'DuPont', 'Working Capital', 'Risk Flags', 'Scenario Analysis']
    assert set(required) <= set(book.sheetnames)
    assert book['Ratios'].freeze_panes == 'B4'
    assert len(book['Income Statement']._charts) == 1
    assert not any(c.data_type == 'f' for ws in book for row in ws for c in row)
    assert book['Executive Summary']['A2'].value.startswith("'")
    for ws in book:
        for row in ws:
            for cell in row:
                assert cell.value not in ['#REF!', '#DIV/0!', '#VALUE!', '#NUM!']


def test_pdf_export():
    import fitz
    a = analyze(demo_company())
    content = pdf_report(a, DEMO_META)
    doc = fitz.open(stream=content, filetype='pdf')
    text = '\n'.join(page.get_text() for page in doc)
    assert 10 <= len(doc) <= 30
    for expected in ['FALCON FINALYSIS', 'Executive summary', 'DuPont analysis', 'Scenario analysis', 'Methodology', 'not investment']:
        assert expected.lower() in text.lower()
    assert 'nan' not in text.lower().split()
    assert all(len(page.get_text().strip()) > 50 for page in doc)


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setenv('FALCON_FINALYSIS_DB', str(tmp_path/'ui.db'))
    at = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30).run()
    next(b for b in at.button if b.label == 'Load Demo Company').click().run()
    at.run()  # Complete the st.switch_page rerun before subsequent interactions.
    assert not at.exception
    assert 'Apex' in at.title[0].value
    assert at.session_state['frame']['Year'].tolist() == [2021, 2022, 2023, 2024, 2025]
    return at


def test_demo_button_opens_overview_from_projects(tmp_path, monkeypatch):
    monkeypatch.setenv('FALCON_FINALYSIS_DB', str(tmp_path/'redirect.db'))
    at = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30).run()
    # AppTest does not render the entrypoint/sidebar after switch_page. Model the
    # post-button state so this still protects the registered-page redirect.
    at.session_state['frame'] = demo_company()
    at.session_state['meta'] = DEMO_META.copy()
    at.session_state['project_id'] = None
    at.session_state['project_name'] = 'Apex • FY2021–FY2025'
    at.session_state['open_demo_overview'] = True
    at.run()
    assert not at.exception
    assert 'Apex Consumer Industries Ltd.' in [title.value for title in at.title]
    assert any(metric.label == 'Revenue' for metric in at.metric)


@pytest.mark.parametrize('page', ['projects', 'market', 'industry_comparison', 'portfolio', 'credgrid', 'valuation', 'data_quality', 'readiness', 'statements', 'ratios', 'dupont', 'trends',
                                  'working_capital', 'cash_flow', 'peers', 'risks', 'scenarios', 'health', 'reports', 'governance', 'about'])
def test_all_pages(app, page):
    app.switch_page(f'pages/{page}.py').run()
    assert not app.exception
    assert not app.error, [e.value for e in app.error]


def test_demo_save_scenario_and_reports(app):
    next(b for b in app.button if b.label == 'Save project').click().run()
    assert app.session_state['project_id']
    app.switch_page('pages/scenarios.py').run()
    next(n for n in app.number_input if n.label == 'Revenue Growth %').set_value(15.).run()
    assert app.session_state['scenario']['revenue_growth'] == .15
    assert not app.error
    next(b for b in app.button if b.label == 'Save scenario').click().run()
    app.switch_page('pages/reports.py').run()
    next(b for b in app.button if b.label == 'Generate PDF & Excel reports').click().run()
    assert app.session_state['pdf_bytes'].startswith(b'%PDF')
    assert app.session_state['xlsx_bytes'].startswith(b'PK')
    assert not app.error


def test_portfolio_demo_calculates_returns_and_covariance(app):
    app.switch_page('pages/portfolio.py').run()
    next(b for b in app.button if b.label == 'Load fictional sample portfolio').click().run()
    next(b for b in app.button if b.label == 'Calculate portfolio return and risk').click().run()
    result = app.session_state['portfolio_analysis']
    assert set(result.asset_summary['Ticker']) == {'FALCON-A', 'FALCON-B', 'FALCON-C'}
    assert result.covariance.shape == (3, 3)
    assert 'Sharpe Ratio' in result.portfolio_summary
    assert any(metric.label == 'Maximum drawdown' for metric in app.metric)
    assert any(button.label == 'Download reviewable portfolio report' for button in app.download_button)
    assert not app.exception and not app.error


def test_credgrid_demo_requires_human_review_and_calculates(app):
    app.switch_page('pages/credgrid.py').run()
    next(button for button in app.button if button.label == 'Load fictional CredGrid case').click().run()
    next(button for button in app.button if button.label == 'Calculate explainable credit analysis').click().run()
    result = app.session_state['cg_analysis']
    assert result.score is not None
    assert result.proposal['Annual Proposed Rate'] > result.proposal['Risk Free Rate']
    assert any(metric.label == 'Modeled amount' for metric in app.metric)
    assert any('trained human' in item.value for item in app.warning)
    assert not app.exception and not app.error


def test_valuation_lab_builds_dcf_and_football_field(app):
    app.switch_page('pages/valuation.py').run()
    next(field for field in app.text_input
         if field.label == 'Market assumptions source and as-of date').set_value(
             'Reviewed demonstration inputs · 2026-09-26').run()
    next(button for button in app.button if button.label == 'Calculate valuation range').click().run()
    assert 'valuation_result' in app.session_state
    assert any(metric.label == 'DCF value per share' for metric in app.metric)
    assert any('Football-field valuation summary' in item.value for item in app.markdown)
    assert not app.exception and not app.error


def test_manual_editor_applies_input(app):
    app.switch_page('pages/projects.py').run()
    app.session_state['statement_editor'] = {'edited_rows': {4: {'Revenue': 1_800_000_000}},
                                             'added_rows': [], 'deleted_rows': []}
    next(b for b in app.button if b.label == 'Validate & apply edits').click().run()
    assert app.session_state['frame'].iloc[-1]['Revenue'] == 1_800_000_000
    assert app.warning  # Reconciliation warnings must remain visible, not repaired.
    assert not app.error


def test_demo_toggle_and_dark_workspace(tmp_path, monkeypatch):
    monkeypatch.setenv('FALCON_FINALYSIS_DB', str(tmp_path/'demo.db'))
    app = AppTest.from_file(str(ROOT/'app.py'), default_timeout=30).run()
    next(t for t in app.toggle if t.label == 'Portfolio Demo Mode').set_value(True).run()
    assert 'Apex' in app.title[0].value
    next(t for t in app.toggle if t.label == 'Dark workspace').set_value(True).run()
    assert not app.exception and not app.error


def test_ten_years_and_twenty_peers():
    import pandas as pd
    early, late = demo_company(), demo_company()
    early['Year'] -= 5
    ten = pd.concat([early, late], ignore_index=True)
    for _ in range(20):
        result = analyze(ten)
        assert len(result.ratios) == 10
    content = pdf_report(result, DEMO_META)
    import fitz
    doc = fitz.open(stream=content, filetype='pdf')
    all_text = '\n'.join(p.get_text() for p in doc)
    assert '2016' in all_text and '2025' in all_text


def test_missing_inputs_remain_unavailable():
    import pandas as pd
    minimal = pd.DataFrame({'Year': [2023, 2024, 2025], 'Revenue': [100., 110., 121.]})
    a = analyze(minimal)
    assert a.health[-1].score is None
    assert a.ratios['Net Profit Margin'].isna().all()
    assert 'N/A' in a.summary['Profitability']
    assert pdf_report(a, DEMO_META).startswith(b'%PDF')
    assert excel_report(a, DEMO_META).startswith(b'PK')


@pytest.mark.parametrize('payload,name', [(b'', 'x.csv'), (b'not a zip', 'x.xlsx'),
                                         (b'a,b', 'x.xls'), (b'Year,Revenue\n', 'x.csv')])
def test_bad_files_have_helpful_errors(payload, name):
    with pytest.raises(ValueError):
        read_file(payload, name)


def test_statement_layout_workbook_extraction_and_scale():
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = 'Statements'
    ws.append(['Amounts in BDT million', 2023, 2024, 2025])
    ws.append(['Revenue', 100, 120, 150])
    ws.append(['Net profit', 10, 11, 14])
    ws.append(['Total assets', 300, 320, 350])
    stream = BytesIO()
    book.save(stream)
    result = extract_workbook(stream.getvalue(), 'annual-report.xlsx')
    assert result.frame.Year.tolist() == [2023, 2024, 2025]
    assert apply_scale(result.frame, 1_000_000).iloc[-1].Revenue == 150_000_000
    assert 'sheet Statements' in result.evidence.Source.iloc[0]


def test_pdf_statement_candidate_extraction():
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    stream = BytesIO()
    story = [Paragraph('Financial statements - amounts in million', getSampleStyleSheet()['Heading1']),
             Table([['Metric', '2023', '2024', '2025'], ['Revenue', '100', '120', '150'],
                    ['Net profit', '10', '11', '14'], ['Total assets', '300', '320', '350']])]
    SimpleDocTemplate(stream, pagesize=A4).build(story)
    result = extract_pdfs([(stream.getvalue(), 'report.pdf')])
    assert result.unit_hint == 'Millions'
    assert result.frame.iloc[-1].Revenue == 150
    assert 'page 1' in result.evidence.Source.iloc[0]


def test_market_history_cleaning_upload_and_workbook():
    raw = pd.DataFrame({'#': [2, 1], 'DATE': ['2025-01-02', '2025-01-01'],
                        'TRADING CODE': ['ABC', 'ABC'], 'LTP*': [11, 10], 'HIGH': [12, 11],
                        'LOW': [10, 9], 'OPENP*': [10, 9], 'CLOSEP*': [11, 10], 'YCP': [10, 9],
                        'TRADE': [20, 10], 'VALUE (mn)': [2.2, 1.0], 'VOLUME': [200, 100]})
    clean = _standardize_dse(raw, 'ABC')
    assert clean.Date.is_monotonic_increasing and clean.Close.tolist() == [10, 11]
    csv = b'Date,Close,Volume\n2025-01-01,10,100\n2025-01-02,11,200\n'
    uploaded = read_price_file(csv, 'prices.csv', 'ABC')
    content = market_data_workbook(uploaded, 'DSE', 'ABC', 'https://www.dsebd.org/example', '2025-01-03T00:00:00Z')
    book = load_workbook(BytesIO(content), data_only=False)
    assert book.sheetnames == ['Summary', 'Price History', 'Monthly Summary', 'Sources']
    assert book['Price History'].freeze_panes == 'A2'
    assert len(book['Summary']._charts) == 1


def test_dse_history_uses_official_legacy_archive(monkeypatch):
    raw = pd.DataFrame({'DATE': ['2026-09-28'], 'TRADING CODE': ['SQURPHARMA'],
                        'OPENP*': [212.0], 'HIGH': [214.0], 'LOW': [211.0],
                        'CLOSEP*': [213.5], 'LTP*': [213.5], 'YCP': [212.0],
                        'TRADE': [100], 'VALUE (MN)': [2.1], 'VOLUME': [10000]})
    seen = {}

    class Response:
        text = '<html></html>'
        url = 'https://old.dsebd.org/day_end_archive.php?example=1'

    def fake_request(method, url, **kwargs):
        seen.update(method=method, url=url, params=kwargs.get('params'))
        return Response()

    monkeypatch.setattr('data.market_data._request', fake_request)
    monkeypatch.setattr('data.market_data._tables', lambda html: [raw])
    frame, source = price_history('DSE', 'SQURPHARMA', date(2026, 9, 1), date(2026, 9, 29))
    assert seen['url'] == 'https://old.dsebd.org/day_end_archive.php'
    assert seen['params']['inst'] == 'SQURPHARMA'
    assert frame.Close.tolist() == [213.5]
    assert source.startswith('https://old.dsebd.org/')


def test_current_dse_company_page_history_is_read_without_estimated_value():
    html = (r'<script>\"series\":[{\"t\":\"28 Sep\",\"date\":\"2026-09-28\",'
            r'\"price\":213.5,\"open\":212.0,\"high\":214.0,\"low\":211.0,'
            r'\"trades\":100,\"volume\":10000}],\"suggestedCode\":\"$undefined\"</script>')
    frame = _modern_dse_history(html, 'SQURPHARMA', date(2026, 9, 1), date(2026, 9, 29))
    assert frame.Close.tolist() == [213.5]
    assert frame.Ticker.tolist() == ['SQURPHARMA']
    assert frame['Value (mn)'].isna().all()


def test_exchange_tls_sessions_are_host_scoped_and_verified():
    for url in ('https://old.dsebd.org/day_end_archive.php',
                'https://www.cse.com.bd/market/marketprice'):
        session = _verified_exchange_session(url)
        adapter = session.get_adapter(url)
        assert adapter._ssl_context.verify_mode == ssl.CERT_REQUIRED
        assert adapter._ssl_context.check_hostname
        assert adapter._ssl_context.verify_flags & ssl.VERIFY_X509_PARTIAL_CHAIN
    assert _verified_exchange_session('https://old.dsebd.org.example.com/') is None


def test_exchange_annual_metrics_prefill_copies_only_direct_net_income():
    table = pd.DataFrame([
        ['Performance Year', 'Net Profit After Tax(mn)', 'Basic EPS Based On', 'Net Asset Value Per Share'],
        ['Performance Year', 'Continuous operations', 'Continuous operations', 'Original'],
        [2023, 100, 5, 40], [2024, 120, 6, 44], [2025, 150, 7, 48],
    ])
    frame, details = _annual_metrics([table], 'https://exchange.example/company/ABC')
    assert frame['Net Income'].tolist() == [100_000_000, 120_000_000, 150_000_000]
    assert frame['Revenue'].isna().all()
    data = exchange_input_template(frame, details, 'DSE', 'ABC',
                                   'https://exchange.example/company/ABC', '2026-01-01T00:00:00Z')
    book = load_workbook(BytesIO(data), data_only=False)
    assert book.sheetnames[:4] == ['Financial Data', 'Exchange Data', 'Exchange Source', 'Field Guide']
    assert book['Financial Data']['L4'].value == 150_000_000  # Net Income column
