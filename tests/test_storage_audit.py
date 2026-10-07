"""Regression checks for private workspace lifecycle, recovery, and safe exports."""
from io import BytesIO
import json
import sqlite3
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from core.accounts import account_path, accounts_enabled
from core.financial_engine import analyze
from data.database import Repository, connect
from data.demo import DEMO_META, demo_company
from data.provenance import (provenance_for_frame, provenance_from_evidence,
                             merge_provenance, changed_provenance)
from data.recovery import project_recovery, read_recovery
from reports.excel_report import exchange_input_template, input_template, market_data_workbook
from reports.pdf_report import pdf_report


def test_workspace_context_commits_and_closes_and_rolls_back(tmp_path):
    path = tmp_path / 'workspace.db'
    with connect(path) as con:
        con.execute('INSERT INTO Companies(company_name) VALUES(?)', ('Committed',))
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        con.execute('SELECT 1')
    with pytest.raises(ValueError, match='abort'):
        with connect(path) as failed:
            failed.execute('INSERT INTO Companies(company_name) VALUES(?)', ('Rolled back',))
            raise ValueError('abort')
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        failed.execute('SELECT 1')
    with connect(path) as read:
        assert [row[0] for row in read.execute('SELECT company_name FROM Companies')] == ['Committed']


def test_failed_schema_initialization_releases_connection(tmp_path, monkeypatch):
    connection = sqlite3.connect(tmp_path / 'invalid.db')
    monkeypatch.setattr('data.database.sqlite3.connect', lambda *args, **kwargs: connection)

    def fail(_con):
        raise sqlite3.DatabaseError('Invalid workspace')

    monkeypatch.setattr('data.database._initialize', fail)
    with pytest.raises(sqlite3.DatabaseError, match='Invalid workspace'):
        connect(tmp_path / 'invalid.db')
    with pytest.raises(sqlite3.ProgrammingError, match='closed'):
        connection.execute('SELECT 1')


def test_backup_includes_committed_wal_content_and_all_workspace_types(tmp_path):
    path = tmp_path / 'workspace.db'
    repo = Repository(path)
    with sqlite3.connect(path) as writer:
        writer.execute('PRAGMA journal_mode=WAL')
        project = repo.save(DEMO_META, demo_company(), 'Financials')
        repo.save_scenario(project, 'Scenario', {'growth': 0.1})
        repo.save_portfolio('Portfolio', {'ticker': 'OLYMPIC'})
        case = repo.save_credit_case('Shop', {'cash': 100})
        repo.save_credit_decision(case, 'Decline', 'Reviewer', 'Insufficient evidence', {})
        repo.save_settings({'tolerance': '0.02'})
        target = tmp_path / 'restored.db'
        target.write_bytes(repo.backup_bytes())
    restored = Repository(target)
    with connect(target) as con:
        assert con.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    assert len(restored.list_projects()) == 1
    assert restored.scenarios(project)[0]['name'] == 'Scenario'
    assert len(restored.list_portfolios()) == 1
    assert len(restored.credit_decisions(case)) == 1
    assert restored.settings() == {'tolerance': '0.02'}


def test_backup_refuses_deleted_workspace(tmp_path):
    repo = Repository(tmp_path / 'missing.db')
    repo.path.unlink()
    with pytest.raises(sqlite3.OperationalError):
        repo.backup_bytes()
    assert not repo.path.exists()


@pytest.mark.parametrize('flag', ['true', ' TRUE ', 'yes', 'on', '1', 'tru', 'unexpected'])
def test_account_flags_fail_closed(monkeypatch, flag):
    monkeypatch.setenv('FALCON_FINALYSIS_ACCOUNTS', flag)
    assert accounts_enabled()


@pytest.mark.parametrize('flag', ['', 'false', ' FALSE ', 'no', 'off', '0'])
def test_explicit_account_disabled_flags(monkeypatch, flag):
    monkeypatch.setenv('FALCON_FINALYSIS_ACCOUNTS', flag)
    assert not accounts_enabled()


def test_account_identity_rejects_boolean_expiry_and_unstructured_claims(tmp_path):
    claims = dict(is_logged_in=True, iss='issuer', sub='alice', exp=True)
    with pytest.raises(ValueError, match='expiry'):
        account_path(str(tmp_path), claims, now=0)
    with pytest.raises(ValueError, match='identity claims'):
        account_path(str(tmp_path), None, now=0)


def test_account_identity_refuses_redirected_workspace(tmp_path, monkeypatch):
    original_resolve = Path.resolve

    def redirected(path, *args, **kwargs):
        if path.name == 'workspace.db':
            return tmp_path / 'another-account' / 'workspace.db'
        return original_resolve(path, *args, **kwargs)

    monkeypatch.setattr(Path, 'resolve', redirected)
    with pytest.raises(ValueError, match='redirected'):
        account_path(str(tmp_path), dict(is_logged_in=True, iss='issuer', sub='alice', exp=500), now=100)


@pytest.mark.parametrize('mutation', [
    lambda doc: doc.update(version=True),
    lambda doc: doc.update(company=[]),
    lambda doc: doc.update(financial_data={'Year': [2021, 2022, 2023]}),
    lambda doc: doc['financial_data'][0].update(Revenue=True),
    lambda doc: doc.update(provenance={}),
    lambda doc: doc['provenance'].append(doc['provenance'][0].copy()),
])
def test_recovery_rejects_malformed_shapes_and_conflicting_lineage(mutation):
    frame = demo_company()
    sources = provenance_for_frame(frame, 'Test', 'Source')
    doc = json.loads(project_recovery(DEMO_META, frame, 'Saved', sources))
    mutation(doc)
    with pytest.raises(ValueError):
        read_recovery(json.dumps(doc).encode())


def test_recovery_rejects_nonstandard_json_and_duplicate_keys():
    payload = project_recovery(DEMO_META, demo_company(), 'Saved')
    with pytest.raises(ValueError, match='JSON numbers'):
        read_recovery(payload.replace(b'"version": 1', b'"version": NaN'))
    with pytest.raises(ValueError, match='duplicate JSON'):
        read_recovery(payload.replace(b'"version": 1', b'"version": 1, "version": 1'))
    with pytest.raises(ValueError, match='valid Falcon'):
        read_recovery(payload.replace(b'"name": "Saved"', b'"name": "\\ud800"'))


def test_market_export_sanitizes_metadata_and_preserves_unknown_volume():
    history = pd.DataFrame({'Date': ['2025-01-03', '2025-01-02'], 'Close': ['12', '10']})
    data = market_data_workbook(history, 'DSE', '=HYPERLINK("bad")',
                                '=HYPERLINK("bad")', '2025-01-03T06:00:00+06:00')
    book = load_workbook(BytesIO(data), data_only=False)
    assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
    assert book['Summary']['B4'].value == 10
    assert book['Summary']['B5'].value == 12
    assert book['Summary']['B7'].value is None
    assert book['Sources']['B5'].value == '03/01/2025 00:00 UTC'
    assert book['Monthly Summary']['E2'].value is None


def test_exchange_and_template_export_escape_headers_sources_and_control_characters():
    frame = demo_company()
    details = pd.DataFrame([[2021, 'Metric\x00', 1, 'BDT', '=HYPERLINK("bad")']],
                           columns=['Year', 'Metric', 'Value', 'Unit', 'Source'])
    output = exchange_input_template(frame, details, 'DSE', '=HYPERLINK("bad")',
                                     '=HYPERLINK("bad")', '03/01/2025')
    book = load_workbook(BytesIO(output), data_only=False)
    assert not any(cell.data_type == 'f' for sheet in book for row in sheet for cell in row)
    assert book['Exchange Data']['B4'].value == 'Metric'
    headers = load_workbook(BytesIO(input_template(pd.DataFrame({'=malicious': [1]}))))
    assert headers['Financial Data']['A1'].data_type != 'f'


def test_pdf_handles_long_source_evidence_and_escapes_markup():
    frame = demo_company()
    reference = '<b>Source evidence</b> ' + 'long evidence ' * 275
    sources = provenance_for_frame(frame, 'Report', reference)
    assert pdf_report(analyze(frame), {**DEMO_META, 'company_name': '<b>Company</b>'},
                      provenance=sources).startswith(b'%PDF')


def test_sources_are_removed_when_values_or_years_are_removed():
    frame = demo_company()
    sources = provenance_for_frame(frame, 'Excel upload', 'Source.xlsx')
    edited = frame.iloc[1:].copy()
    edited.loc[edited.index[-1], 'Revenue'] = float('nan')
    merged = merge_provenance(sources, changed_provenance(frame, edited), accepted_frame=edited)
    assert int(frame.Year.min()) not in merged.Year.tolist()
    assert not ((merged.Year == int(edited.Year.max())) & (merged.Field == 'Revenue')).any()


def test_reviewed_extraction_records_corrections_added_values_and_unit_conversion():
    evidence = pd.DataFrame([
        {'Year': 2025, 'Falcon Finalysis field': 'Revenue', 'Value': 10, 'Source': 'report.pdf, page 1'},
        {'Year': 2025, 'Falcon Finalysis field': 'Net Income', 'Value': 2, 'Source': 'report.pdf, page 2'},
    ])
    reviewed = pd.DataFrame([{'Year': 2025, 'Revenue': 11_000_000, 'Net Income': 2_000_000,
                              'Cash': 1_000_000}])
    sources = provenance_from_evidence(evidence, accepted_frame=reviewed,
                                       multiplier=1_000_000).set_index('Field')
    assert sources.at['Revenue', 'Source Type'] == 'Manual entry'
    assert 'report.pdf, page 1' in sources.at['Revenue', 'Source Reference']
    assert sources.at['Cash', 'Source Type'] == 'Manual entry'
    assert sources.at['Net Income', 'Source Type'] == 'Annual report extraction'
    assert 'unit multiplier: 1e+06' in sources.at['Net Income', 'Source Reference']
